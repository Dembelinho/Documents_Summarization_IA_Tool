import unittest
import os
from docx import Document as DocxDocument
from fpdf import FPDF
from app.document_processor import DocumentProcessor, load_document, SummaryPrompt
from langchain.schema import Document
from langchain_core.runnables import Runnable


class TestDocumentProcessor(unittest.TestCase):
    def setUp(self):
        # Set up a test file directory
        self.test_files_dir = os.path.join(os.path.dirname(__file__), "test_files")
        os.makedirs(self.test_files_dir, exist_ok=True)

        # Create a valid PDF file
        self.pdf_file = os.path.join(self.test_files_dir, "test.pdf")
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt="This is a test PDF file.", ln=True, align="C")
        pdf.output(self.pdf_file)

        # Create a valid TXT file
        self.txt_file = os.path.join(self.test_files_dir, "test.txt")
        with open(self.txt_file, "w") as f:
            f.write("This is a test TXT file.")

        # Create a valid DOCX file
        self.docx_file = os.path.join(self.test_files_dir, "test.docx")
        doc = DocxDocument()
        doc.add_paragraph("This is a test DOCX file.")
        doc.save(self.docx_file)

        # Initialize DocumentProcessor with a mock API key
        self.processor = DocumentProcessor(api_key="mock-api-key")

    def tearDown(self):
        # Clean up test files
        for file in [self.pdf_file, self.txt_file, self.docx_file]:
            if os.path.exists(file):
                os.remove(file)

    def test_load_document_pdf(self):
        documents = load_document(self.pdf_file)
        self.assertIsInstance(documents, list)
        self.assertGreater(len(documents), 0)
        self.assertIsInstance(documents[0], Document)

    def test_load_document_txt(self):
        documents = load_document(self.txt_file)
        self.assertIsInstance(documents, list)
        self.assertGreater(len(documents), 0)
        self.assertIsInstance(documents[0], Document)

    def test_load_document_docx(self):
        documents = load_document(self.docx_file)
        self.assertIsInstance(documents, list)
        self.assertGreater(len(documents), 0)
        self.assertIsInstance(documents[0], Document)

    def test_load_document_unsupported_format(self):
        with self.assertRaises(ValueError):
            load_document(os.path.join(self.test_files_dir, "test.unsupported"))

    def test_summary_prompt_english(self):
        prompt = SummaryPrompt.get_prompt("default", "concise", "english")
        self.assertIn("Write a very concise summary in English", prompt.template)

    def test_summary_prompt_french(self):
        prompt = SummaryPrompt.get_prompt("default", "concise", "french")
        self.assertIn("Écrivez un résumé très concis en français", prompt.template)

    def test_process_document(self):
        # Mock the summarization chain
        class MockRunnable(Runnable):
            def invoke(self, input):
                return "Mock summary"

        self.processor.llm = MockRunnable()

        summary = self.processor.process_document(self.txt_file, summary_type="default", length="concise", language="english")
        self.assertEqual(summary, "Mock summary")

if __name__ == "__main__":
    unittest.main()